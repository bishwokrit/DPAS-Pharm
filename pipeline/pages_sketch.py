"""Sketch mode (hidden): Kahoot spreadsheets built from each unit's practice-exam questions, and a map of
Sketchy Pharm videos to our lectures and drugs. The tab only appears after Ctrl+Alt+S and the password.

Sketchy itself is paid, copyrighted material: the raw docs stay in the private repo (dpas-textbook/sketchy/).
Only work/sketchy.json is published — video titles and which of our drugs/lectures they cover, plus notes
written in our own words — never Sketchy's scenes, symbols or text."""
import hashlib, html, json, re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
TEMPLATE = ROOT / "pipeline" / "templates" / "kahoot_template.xlsx"
SHORT = json.loads((ROOT / "pipeline" / "kahoot_stems.json").read_text())  # hand-shortened stems (Kahoot caps at 120 chars)
DECK_MAX = 300  # questions per lecture deck


def _h(s):
    return int(hashlib.sha1(s.encode()).hexdigest(), 16)


def _plain(t):
    return re.sub(r"\s+", " ", t.replace("**", "")).strip()


def _secs(stem, texts):
    calc = bool(re.search(r"\d", stem)) and len(stem) >= 90
    return 60 if calc else 20 if len(stem) < 80 and max(map(len, texts)) < 40 else 30


def kahoot_rows(q):
    """One multiple-choice question as a Kahoot row: stem ≤120, the answer plus 3 distractors (≤75 each), shuffled."""
    stem = _plain(SHORT.get(q["id"], q["stem"]))
    wrong = sorted((k for k in q["options"] if k != q["answer"]), key=lambda k: _h(q["id"] + k))[:3]
    opts = sorted([q["answer"], *wrong], key=lambda k: _h(q["id"] + "pos" + k))
    texts = [_plain(q["options"][k]) for k in opts]
    if len(stem) > 120 or any(len(t) > 75 for t in texts) or len(set(texts)) < 4:
        return None
    return [stem, *texts, _secs(stem, texts), str(opts.index(q["answer"]) + 1)]


def card_rows(cards, lid):
    """Flashcards as quick-fire rounds: the card's front, its own answer, and 3 other answers from the same lecture
    (closest in length, so the right one doesn't stand out)."""
    ok = [c for c in cards if len(_plain(c["a"])) <= 75 and len(_plain(c["q"])) <= 100]
    rows = []
    for c in ok:
        a = _plain(c["a"])
        others = sorted({_plain(o["a"]) for o in ok if _plain(o["a"]) != a}, key=lambda t: (abs(len(t) - len(a)), _h(lid + a + t)))[:3]
        if len(others) < 3:
            continue
        opts = sorted([a, *others], key=lambda t: _h(lid + c["q"] + t))
        stem = f"{_plain(c['q'])}: which is right?"
        rows.append([stem, *opts, 20 if len(a) < 40 else 30, str(opts.index(a) + 1)])
    return rows


def _write(rows, path):
    """Fill the Kahoot template (row 8 is its header); skip the write when nothing changed (xlsx stores a save time)."""
    if path.exists():
        old = openpyxl.load_workbook(path)["Sheet1"]
        if [[old[f"{c}{9 + i}"].value for c in "BCDEFGH"] for i in range(len(rows) + 1)] == [*rows, [None] * 7]:
            return
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb["Sheet1"]
    for i in range(9, ws.max_row + 1):
        for col in "ABCDEFGH":
            ws[f"{col}{i}"] = None
    for i, r in enumerate(rows):
        ws[f"A{9 + i}"] = i + 1
        for col, v in zip("BCDEFGH", r):
            ws[f"{col}{9 + i}"] = v
    wb.save(path)


def build_kahoots():
    """docs/sketch/kahoot/unit<N>_<lecture>.xlsx: one deck per lecture (practice-exam questions written for that
    lecture, its self-test questions that fit Kahoot's limits, then its flashcards), grouped by unit."""
    out_dir = DOCS / "sketch" / "kahoot"
    out_dir.mkdir(parents=True, exist_ok=True)
    syl = json.loads((ROOT / "work" / "syllabus.json").read_text())
    exam_qs = []
    for meta in sorted((DOCS / "exams").glob("*/meta.json")):
        exam_qs += json.loads((meta.parent / "quiz.json").read_text())
    keep, units, skipped, all_rows = set(), {}, 0, {}
    for lec in syl["lectures"]:
        lid, kbp = lec["id"], DOCS / lec["id"] / "kb.json"
        if not kbp.exists():
            continue
        kb = json.loads(kbp.read_text())
        mine = [q for q in exam_qs if ((q.get("link") or {}).get("href", "").split("/")[2:3] == [lid]) or q.get("src", "").startswith(lid + " ")]
        rows, seen = [], set()
        for q in sorted(mine, key=lambda q: _h("mix" + q["id"])) + [{**q, "id": f"{lid}:{q['id']}"} for q in kb["quiz"]]:
            r = kahoot_rows(q)
            if r and r[0] not in seen:
                rows.append(r); seen.add(r[0])
            skipped += r is None
        rows += [r for r in card_rows(kb["cards"], lid) if r[0] not in seen]
        rows = rows[:DECK_MAX]
        all_rows[lid] = rows
        name = f"unit{lec['unit']}_{lid}.xlsx"
        keep.add(name)
        _write(rows, out_dir / name)
        units.setdefault(lec["unit"], []).append({"lid": lid, "title": kb["title"], "name": name, "n": len(rows)})
    keep.add("template.xlsx")
    for old in out_dir.glob("*.xlsx"):
        (old.unlink() if old.name not in keep else None)
    # the game-night builder mixes these rows in the browser and fills the same template
    if not (out_dir / "template.xlsx").exists() or (out_dir / "template.xlsx").read_bytes() != TEMPLATE.read_bytes():
        (out_dir / "template.xlsx").write_bytes(TEMPLATE.read_bytes())
    (out_dir / "rows.json").write_text(json.dumps(all_rows, ensure_ascii=False))
    if skipped:
        print(f"  sketch: {skipped} questions too long for Kahoot (add a short stem in pipeline/kahoot_stems.json)")
    return [{"unit": u, "name": syl["units"].get(str(u), ""), "decks": d, "total": sum(x["n"] for x in d)} for u, d in sorted(units.items())]


def sketch_page(page, SITE_NAME):
    kahoots = build_kahoots()
    sk_file = ROOT / "work" / "sketchy.json"
    sk = json.loads(sk_file.read_text()) if sk_file.exists() else {"videos": []}
    cards = "".join(
        f'<div class="sx-unit"><div class="sx-unit-h"><strong>Unit {u["unit"]} <span>· {html.escape(u["name"])}</span></strong><span>{u["total"]} questions · {len(u["decks"])} deck{"s" if len(u["decks"]) != 1 else ""}</span></div>'
        f'<div class="sx-files">' + "".join(
            f'<a class="sx-file" href="kahoot/{d["name"]}" download><span>{d["lid"]}</span><em>{html.escape(d["title"])}</em><small>{d["n"]} questions · .xlsx</small></a>'
            for d in u["decks"]) + "</div></div>" for u in kahoots)
    vids = sk.get("videos", [])
    if vids:
        sketchy = '<div id="tracker-app"></div>'
    else:
        sketchy = '<div class="sx-empty"><strong>Waiting for your Sketchy docs</strong><span>Once they are uploaded, every Sketchy Pharm video is matched to the lectures and drugs it covers, right here.</span></div>'
    import sys
    sys.path.insert(0, str(ROOT / "work"))
    import drugs as DR
    syl = json.loads((ROOT / "work" / "syllabus.json").read_text())
    built = {l["id"] for l in syl["lectures"] if (DOCS / l["id"] / "kb.json").exists()}
    cfg = {"videos": vids, "units": syl["units"], "exams": syl.get("exams", []),
           "lectures": [{"id": l["id"], "unit": l["unit"], "date": l["date"], "title": json.loads((DOCS / l["id"] / "kb.json").read_text())["title"]} for l in syl["lectures"] if l["id"] in built],
           "quizzes": [p.parent.name for p in sorted((DOCS / "exams").glob("*/quiz.json"))],
           "drugs": [{"name": d["name"], "unit": d["unit"], "cls": d["cls"], "pattern": "|".join(re.escape(n) for n in [d["name"], *d["aliases"]]),
                      "facts": [{"text": f["text"], "src": f["src"]} for f in d["facts"]]} for d in DR.DRUGS]}
    cfg_json = json.dumps(cfg, ensure_ascii=False).replace("</", "<\\/")
    sec = lambda i, k, kick, h, inner="": f'<section class="section sx-sec" id="{i}"><div class="wrap"><p class="kicker">{k} · {kick}</p><h2>{h}</h2>{inner}<div id="{i}-app"></div></div></section>'
    body = f"""<header class="hero wrap"><p class="eyebrow">Sketch</p><h1>Sketch mode</h1>
<p class="lede">Sketchy Pharm matched to the lectures, your own memory hooks, Kahoot decks and games. Only you can see this tab.</p>
<nav class="sx-jump"><a href="#sotd">Today</a><a href="#sketchy">Sketchy</a><a href="#recall">Recall</a><a href="#hooks">Hooks</a><a href="#kahoot">Kahoot</a><a href="#h2h">Head-to-head</a></nav>
<p><button class="btn ghost small" id="sx-lock" type="button">Lock Sketch mode</button></p></header>
{sec("sotd", 1, "Sketch of the day", "One drug, its sketch, three questions")}
{sec("sketchy", 2, "Sketchy Pharm", "Watch the right sketch", f'<div id="playlist-app"></div>{sketchy}')}
{sec("recall", 3, "Recall drill", "Say it before you see it")}
{sec("hooks", 4, "Memory hooks", "Your words, on every drug sheet")}
<section class="section sx-sec" id="kahoot"><div class="wrap"><p class="kicker">5 · Kahoot</p><h2>A deck for every lecture</h2>
<p class="sx-how">In Kahoot: <b>Create</b> → <b>Kahoot</b> → <b>Add question</b> → <b>Import spreadsheet</b>, then drop in a deck. Each deck holds that lecture's practice-exam and self-test questions, then quick-fire flashcard rounds.</p>
<div class="sx-units">{cards}</div><h3 class="sx-h3">Game-night builder</h3><div id="kahoot-app"></div></div></section>
{sec("h2h", 6, "Head-to-head", "Two players, one screen, ten questions")}
<script type="application/json" id="sketch-data">{cfg_json}</script>
<script>if(!(function(){{try{{return JSON.parse(localStorage.getItem("sketch"))}}catch(e){{}}}})())location.replace("../index.html");
document.getElementById("sx-lock").onclick=function(){{window.dpasSketchLock?window.dpasSketchLock():(localStorage.removeItem("sketch"),location.href="../index.html")}};</script>
<script src="../assets/sketch.js"></script>"""
    out = DOCS / "sketch"
    out.mkdir(exist_ok=True)
    (out / "index.html").write_text(page(f"Sketch · {SITE_NAME}", body, depth=1, active="sketch"))
    print(f"  sketch: {sum(len(u['decks']) for u in kahoots)} Kahoot decks ({sum(u['total'] for u in kahoots)} questions), {len(vids)} Sketchy videos mapped")
