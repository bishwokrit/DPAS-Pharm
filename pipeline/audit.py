"""Coverage audit: is every fact from the slides (and textbook pages) in the deck + NotebookLM pack?

Checks
  1. Every slide line / table row / speaker note, every textbook sentence, and every study-guide
     item assigned to this lecture is matched by content citing that source (word-overlap match).
  2. Slide emphasis (bold/highlight/colored) shows up bold in the content.
  3. Every fact actually made it into the .pptx and the NotebookLM source.
  4. Practice questions follow the rules (A–E balance, no letter runs, similar option lengths, ~50/50 mix).
  5. Objectives with thin coverage.
Writes outputs/<LEC>/<LEC>_audit.md and work/<LEC>/audit.json. Exit code 2 = something missing.
"""
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from pptx import Presentation

from check_layout import check as check_layout
from content import sections_for

STOP = set("""a an the and or of to in on for with by is are was were be been being as at from that this these those it its
into than then there their which who whom whose what when where why how not no but if so such can may might should would will
also more most less least very each per via other others about after before during between over under within without both
either neither all any some one two three use used using""".split())


def plainify(text):
    # NFKC turns subscripts/superscripts into plain digits (TD₅₀ -> TD50)
    return unicodedata.normalize("NFKC", str(text)).lower().replace("**", "")


def toks(text):
    out = set()
    for w in re.findall(r"[a-z0-9]+", plainify(text)):
        if w in STOP or (len(w) < 2 and not w.isdigit()):
            continue
        out.add(w[:-1] if len(w) > 4 and w.endswith("s") else w)
    return out


def norm(text):
    return re.sub(r"[^a-z0-9]+", "", plainify(text))


def cited_refs(src):
    """'S12; S14-16; p. 245-246; SG U1-R3' -> {'S12','S14','S15','S16','p245','p246','SG:U1-R3'}"""
    refs = {f"SG:{m}" for m in re.findall(r"SG\s*(U\d+-[A-Z]\d*)", src or "")}
    src = re.sub(r"SG\s*U\d+-[A-Z]\d*", "", src or "")
    for a, b in re.findall(r"S\s*(\d+)\s*(?:[-–]\s*S?\s*(\d+))?", src or ""):
        for n in range(int(a), int(b or a) + 1):
            refs.add(f"S{n}")
    for a, b in re.findall(r"pp?\.\s*(\d+)\s*(?:[-–]\s*(\d+))?", src or ""):
        for n in range(int(a), int(b or a) + 1):
            refs.add(f"p{n}")
    return refs


def content_items(c):
    """Every piece of text in content.json with its citation."""
    items = [{"text": f["text"], "src": f.get("src", ""), "where": "About this lecture", "fact": True} for f in c.get("about", [])]
    items += [{"text": o["text"], "src": o.get("src", ""), "where": "Objectives", "fact": False} for o in c.get("objectives", [])]
    for cl in c["classes"]:
        for fig in cl.get("figures", []):
            items.append({"text": fig["caption"], "src": fig.get("src", ""), "where": f"{cl['name']} / figure", "fact": True})
        for key, _ in sections_for(cl):
            for f in cl["sections"].get(key, []):
                items.append({"text": f["text"], "src": f.get("src", ""), "where": f"{cl['name']} / {key}", "fact": True})
        for f in cl.get("high_yield", []):
            items.append({"text": f["text"], "src": f.get("src", ""), "where": f"{cl['name']} / high-yield", "fact": True})
        for h in cl.get("hooks", []):
            items.append({"text": h["text"], "src": h.get("src", ""), "where": f"{cl['name']} / hook", "fact": True})
        for t in cl.get("comparisons", []):
            for r in t["rows"]:
                items.append({"text": " ".join(map(str, r)) + " " + " ".join(t["columns"]), "src": t.get("src", ""), "where": f"{cl['name']} / table", "fact": False})
        for cf in cl.get("conflicts", []):
            for side in ("slides", "textbook"):
                items.append({"text": cf[side]["text"], "src": cf[side]["src"], "where": f"{cl['name']} / conflict", "fact": True})
        for q in cl.get("questions", []):
            blob = " ".join([q["stem"], q["correct"], q.get("why_correct", "")] + [d["text"] + " " + d["why"] for d in q["distractors"]])
            # a lecture quick check rebuilt as a practice question covers its own slides
            items.append({"text": blob, "src": q.get("from_slides", "") or q.get("src", ""), "where": f"{cl['name']} / question",
                          "fact": False, "question": not q.get("from_slides")})
    return items


def source_units(slides, textbook, guide):
    units = []
    for s in slides:
        ref = f"S{s['slide']}"
        divider = any(it["text"].strip().upper() == "SECTION" for it in s["items"])
        for it in s["items"]:
            if it.get("heading") or divider:
                continue  # titles, labels, and section-divider taglines carry no facts
            units.append({"ref": ref, "text": it["text"], "kind": "slide"})
        for t in s["tables"]:
            head = t[0] if t else []
            for row in t[1:] or t:
                units.append({"ref": ref, "text": " ".join(row) + " " + " ".join(head), "kind": "slide table"})
        for sent in re.split(r"(?<=[.!?])\s+", s.get("notes", "")):
            units.append({"ref": ref, "text": sent, "kind": "speaker note"})
    for p in textbook:
        page = re.sub(r"\D", "", p["page"]) or p["page"]
        for sent in re.split(r"(?<=[.!?])\s+", p["text"].replace("\n", " ")):
            if len(sent.split()) >= 6:
                units.append({"ref": f"p{page}", "text": sent, "kind": "textbook"})
    for gid, g in guide.items():
        # matching exercise: the definitions list is the content; the instructions aren't
        text = g["question"].rsplit("Column B", 1)[-1] if gid.endswith("-M") else g["answer"]
        for sent in re.split(r"(?<=[.;])\s+|\s\(\d+\)\s", text):
            if len(sent.split()) >= 4:
                units.append({"ref": f"SG:{gid}", "text": sent, "kind": "study guide"})
    return [u for u in units if len(toks(u["text"])) >= 2]


def text_of_pptx(path):
    chunks = []
    for slide in Presentation(path).slides:
        for sh in slide.shapes:
            if sh.has_text_frame:
                chunks.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                chunks += [cell.text for row in sh.table.rows for cell in row.cells]
    return norm(" ".join(chunks))


def question_checks(c):
    qs = [q for cl in c["classes"] for q in cl.get("questions", [])]
    issues = []
    if not qs:
        return ["No practice questions."], {}
    letters = [q["answer"] for q in qs]
    dist = Counter(letters)
    for i in range(len(letters) - 2):
        if letters[i] == letters[i + 1] == letters[i + 2]:
            issues.append(f"Answer letter {letters[i]} three times in a row (Q{i + 1}–Q{i + 3}).")
    kinds = Counter(q["kind"] for q in qs)
    share = kinds["vignette"] / len(qs)
    if len(qs) >= 4 and not 0.35 <= share <= 0.65:
        issues.append(f"Vignette share is {share:.0%} (target ~50%).")
    for q in qs:
        lens = {L: len(o["text"]) for L, o in q["options"].items()}
        correct_len = lens[q["answer"]]
        others = [v for L, v in lens.items() if L != q["answer"]]
        if correct_len > 1.3 * max(others):
            issues.append(f"Q{q['number']}: correct answer is noticeably the longest option.")
        if max(lens.values()) > 2.2 * max(1, min(lens.values())):
            issues.append(f"Q{q['number']}: option lengths very uneven ({min(lens.values())}–{max(lens.values())} chars).")
    return issues, {"count": len(qs), "letters": dict(sorted(dist.items())), "vignette": kinds["vignette"], "mechanism": kinds["mechanism"]}


def main(lec, work_dir, out_dir, deck_path, nlm_dir, threshold, guide_ids=()):
    threshold = float(threshold)
    work, out = Path(work_dir), Path(out_dir)
    c = json.loads((work / "content.resolved.json").read_text())
    slides = json.loads((work / "slides.json").read_text())
    tb_path = work / "textbook.json"
    textbook = json.loads(tb_path.read_text()) if tb_path.exists() else []
    sg_path = work.parent / "study_guide.json"
    all_guide = json.loads(sg_path.read_text()) if sg_path.exists() else {}
    guide = {k: all_guide[k] for k in guide_ids if k in all_guide}

    items = content_items(c)
    # questions don't count as coverage: a fact must be taught, not just asked about
    teach = [i for i in items if not i.get("question")]
    by_ref = {}
    for it in teach:
        for r in cited_refs(it["src"]):
            by_ref.setdefault(r, set()).update(toks(it["text"]))
    all_toks = set().union(*(toks(i["text"]) for i in teach)) if teach else set()

    missing, uncited, skipped = [], [], []
    units = source_units(slides, textbook, guide)
    skips = c.get("audit_skip", [])
    for u in units:
        sk = next((k for k in skips if k["ref"] == u["ref"] and norm(u["text"]).startswith(norm(k["starts_with"]))), None)
        if sk:
            skipped.append({**u, "reason": sk["reason"]})
            continue
        ut = toks(u["text"])
        cited_score = len(ut & by_ref.get(u["ref"], set())) / len(ut)
        if cited_score >= threshold:
            continue
        if len(ut & all_toks) / len(ut) >= threshold:
            uncited.append(u)
        else:
            missing.append({**u, "score": round(cited_score, 2)})

    # emphasis on slides -> bold in content
    bold_text = norm(" ".join(m for it in items for m in re.findall(r"\*\*(.+?)\*\*", it["text"])))
    emph_missing = []
    for s in slides:
        # quiz/answer slides and the title slide style everything bold; that isn't emphasis
        if s["slide"] == 1 or re.match(r"(quick check|answer)\b", s.get("title", ""), re.I):
            continue
        for it in s["items"]:
            for e in it["emphasis"]:
                strong = {"bold", "underline", "highlight"} & set(e["marks"])
                label = norm(e["text"]) in ("why", "section", "askyourself", "clinicalpearl")
                if strong and not label and len(norm(e["text"])) >= 3 and norm(e["text"]) not in bold_text:
                    emph_missing.append({"ref": f"S{s['slide']}", "text": e["text"], "marks": e["marks"]})

    # did every fact reach the deck and the NotebookLM pack?
    deck_text = text_of_pptx(deck_path)
    nlm_text = norm(" ".join(p.read_text() for p in Path(nlm_dir).glob("*_source*.md")))
    not_in_deck = [i for i in items if i["fact"] and norm(i["text"]) not in deck_text]
    not_in_nlm = [i for i in items if i["fact"] and norm(i["text"]) not in nlm_text]

    q_issues, q_stats = question_checks(c)
    layout = check_layout(deck_path)

    thin = []
    for o in c["objectives"]:
        cls = [cl for cl in c["classes"] if o["id"] in cl.get("objectives", [])]
        n = sum(len(f) for cl in cls for f in cl["sections"].values())
        if not cls or n < 4 or o.get("thin_note"):
            thin.append(f"{o['id']}: {o['text']} — {len(cls)} class(es), {n} facts" + (f" ({o['thin_note']})" if o.get("thin_note") else ""))

    passed = not (missing or not_in_deck or not_in_nlm or layout)
    result = {"lecture": lec, "passed": passed, "units_checked": len(units), "missing": missing, "uncited": uncited,
              "emphasis_not_bold": emph_missing, "skipped": skipped, "not_in_deck": not_in_deck, "not_in_notebooklm": not_in_nlm,
              "layout_problems": layout, "question_issues": q_issues, "question_stats": q_stats, "thin_objectives": thin}
    (work / "audit.json").write_text(json.dumps(result, indent=1, ensure_ascii=False))

    def lst(rows, fmt):
        return "\n".join(fmt(r) for r in rows) if rows else "- None ✅"

    tb_note = f"{len(textbook)} textbook pages checked." if textbook else "Textbook pending — slides only."
    report = f"""# Coverage audit — {lec}: {c['lecture']['title']}

**Result: {'PASS ✅ nothing missing' if passed else 'NEEDS FIXES ❌'}**

Checked {len(units)} source lines ({len(slides)} slides; {len(guide)} study-guide items; {tb_note})

## Missing from the study materials
{lst(missing, lambda u: f"- {u['ref']} ({u['kind']}): {u['text'][:160]}")}

## Covered but cited to a different slide/page (check citation)
{lst(uncited, lambda u: f"- {u['ref']}: {u['text'][:160]}")}

## Left out on purpose (no study content)
{lst(skipped, lambda u: f"- {u['ref']}: {u['text'][:110]} — _{u['reason']}_")}

## Emphasized on slides but not bold in deck
{lst(emph_missing, lambda e: f"- {e['ref']}: \"{e['text']}\" ({', '.join(e['marks'])})")}

## Facts that did not reach the deck / NotebookLM
{lst(not_in_deck, lambda i: f"- Deck: {i['where']}: {i['text'][:120]}")}
{lst(not_in_nlm, lambda i: f"- NotebookLM: {i['where']}: {i['text'][:120]}") if not_in_nlm else ''}

## Deck layout (text overflow / overlap / off-slide)
{lst(layout, lambda s: f"- {s}")}

## Practice questions
{q_stats.get('count', 0)} questions · {q_stats.get('vignette', 0)} vignette / {q_stats.get('mechanism', 0)} mechanism · answer letters {q_stats.get('letters', {})}
{lst(q_issues, lambda s: f"- {s}")}

## Objectives with thin coverage
{lst(thin, lambda s: f"- ⚠ {s}")}
"""
    (out / f"{lec}_audit.md").write_text(report)
    print(f"  audit: {'PASS' if passed else 'NEEDS FIXES'} — {len(missing)} missing, {len(uncited)} citation checks, "
          f"{len(emph_missing)} emphasis, {len(not_in_deck)} not in deck, {len(layout)} layout, {len(q_issues)} question notes")
    return 0 if passed else 2


if __name__ == "__main__":
    sys.exit(main(*sys.argv[1:7], guide_ids=sys.argv[7].split(",") if len(sys.argv) > 7 else ()))
