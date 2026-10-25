"""Check a finished deck for layout problems the way PowerPoint will show them.

For every text box: does the text fit (word-wrapped with real Calibri widths)? Does it overlap
another text box? Does it run off the slide? Returns a list of problems (empty = clean).
"""
import sys
from functools import lru_cache

from PIL import ImageFont
from pptx import Presentation

NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FONT = {False: "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
        True: "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"}
IN = 914400


@lru_cache(maxsize=None)
def font(bold):
    return ImageFont.truetype(FONT[bold], 1000)


def width_in(text, pt, bold):
    return font(bold).getlength(text) / 1000 * pt / 72


def para_height(p, box_w, default_pt):
    """Height (inches) of one paragraph after word wrap."""
    ppr = p._p.find(f"{NS}pPr")
    mar_l = int(ppr.get("marL", 0)) / IN if ppr is not None else 0
    avail = box_w - mar_l
    words, max_pt = [], 0
    for r in p.runs:
        pt = r.font.size.pt if r.font.size else default_pt
        max_pt = max(max_pt, pt)
        for w in r.text.split(" "):
            words.append((w + " ", pt, bool(r.font.bold)))
    if not words:
        max_pt = max_pt or default_pt
        return max_pt * 1.2 / 72
    n, cur = 1, 0.0
    for w, pt, bold in words:
        ww = width_in(w, pt, bold)
        if cur > 0 and cur + width_in(w.rstrip(), pt, bold) > avail:
            n, cur = n + 1, 0.0
        cur += ww
    after = 0
    if ppr is not None and ppr.find(f"{NS}spcAft/{NS}spcPts") is not None:
        after = int(ppr.find(f"{NS}spcAft/{NS}spcPts").get("val")) / 100 / 72
    return n * max_pt * 1.2 / 72 + after


def insets(shape):
    body = shape.text_frame._txBody.find(f"{NS}bodyPr")
    get = lambda k, d: int(body.get(k, d)) / IN if body is not None else d / IN
    return get("lIns", 91440), get("tIns", 45720), get("rIns", 91440), get("bIns", 45720)


def check(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width / IN, prs.slide_height / IN
    problems = []
    for n, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shp in slide.shapes:
            if not shp.has_text_frame or not shp.text_frame.text.strip():
                continue
            x, y, w, h = shp.left / IN, shp.top / IN, shp.width / IN, shp.height / IN
            l, t, r, b = insets(shp)
            need = sum(para_height(p, w - l - r, 18) for p in shp.text_frame.paragraphs) + t + b
            label = shp.text_frame.text.strip().replace("\n", " ")[:50]
            if need > h + 0.06:
                problems.append(f"slide {n}: text overflows its box by {need - h:.2f}\" — \"{label}\"")
            if x < -0.01 or y < -0.01 or x + w > sw + 0.01 or y + max(h, need) > sh + 0.01:
                problems.append(f"slide {n}: text runs off the slide — \"{label}\"")
            boxes.append((x, y, w, max(h, need), label))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                ax, ay, aw, ah, al = boxes[i]
                bx, by, bw, bh, bl = boxes[j]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.05 and oy > 0.05:
                    problems.append(f"slide {n}: text boxes overlap — \"{al[:30]}\" / \"{bl[:30]}\"")
    return problems


if __name__ == "__main__":
    probs = check(sys.argv[1])
    print("\n".join(probs) or "layout clean")
