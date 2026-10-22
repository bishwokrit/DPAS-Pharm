"""Pull every piece of text out of a lecture's .pptx, keeping slide numbers and emphasis.

Writes work/<LEC>/slides.json (for the audit) and slides.md (easy to read while authoring).
Emphasis = bold, underline, highlight, or a text color different from the deck's normal body text.
Headings (big title text, small ALL-CAPS labels, white text on colored boxes) are not emphasis.
"""
import json
import sys
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
PLAIN_COLORS = {None, "000000", "FFFFFF", "262626", "404040", "595959", "1F1F1F", "333333"}


def run_color(run):
    try:
        c = run.font.color
        return str(c.rgb) if c and c.type is not None and c.rgb is not None else None
    except (AttributeError, TypeError):
        return None


def run_size(run):
    return run.font.size.pt if run.font.size else None


def is_heading_run(run):
    size, text = run_size(run), run.text.strip()
    return (size and size >= 20) or (size and size <= 12 and text.isupper()) or run_color(run) == "FFFFFF"


def run_emphasis(run, plain_colors):
    marks = []
    f = run.font
    if f.bold:
        marks.append("bold")
    if f.underline:
        marks.append("underline")
    if run._r.find(f"{NS_A}rPr/{NS_A}highlight") is not None:
        marks.append("highlight")
    if run_color(run) not in plain_colors:
        marks.append("color")
    return marks


def paragraph_record(p, plain_colors):
    runs = [r for r in p.runs if r.text]
    text = "".join(r.text for r in runs).strip()
    if not text:
        return None
    heading = all(is_heading_run(r) for r in runs if r.text.strip())
    md, emph = [], []
    for r in runs:
        marks = [] if heading else run_emphasis(r, plain_colors)
        if marks and r.text.strip():
            emph.append({"text": r.text.strip(), "marks": marks})
            md.append(f"**{r.text.strip()}** ")
        else:
            md.append(r.text)
    size = max((run_size(r) or 0) for r in runs)
    return {"text": text, "level": p.level, "emphasis": emph, "md": "".join(md).strip(), "heading": heading, "size": size}


def body_colors(prs):
    """The deck's normal body-text color(s): the most common color in regular-size, non-bold text."""
    count = Counter()
    for slide in prs.slides:
        for s in walk_shapes(slide.shapes):
            if getattr(s, "has_text_frame", False) and s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip() and not r.font.bold and not is_heading_run(r):
                            count[run_color(r)] += len(r.text)
    common = [c for c, _ in count.most_common(1)]
    return PLAIN_COLORS | set(common)


def walk_shapes(shapes):
    """Yield shapes in reading order (top-to-bottom, then left-to-right), opening groups."""
    def key(s):
        return (round((s.top or 0) / Emu(300000)), s.left or 0)
    for s in sorted(shapes, key=key):
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(s.shapes)
        else:
            yield s


def extract(pptx_path, out_dir=None):
    prs = Presentation(pptx_path)
    slide_area = prs.slide_width * prs.slide_height
    plain = body_colors(prs)
    slides = []
    for i, slide in enumerate(prs.slides, start=1):
        rec = {"slide": i, "title": "", "items": [], "tables": [], "images": 0, "figures": [], "notes": ""}
        title_shape = slide.shapes.title
        if title_shape is not None and title_shape.has_text_frame:
            rec["title"] = title_shape.text_frame.text.strip()
        for s in walk_shapes(slide.shapes):
            if title_shape is not None and s.shape_id == title_shape.shape_id:
                continue
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                rec["images"] += 1
                # big pictures are figures worth reusing; small ones are icons
                if out_dir and s.width * s.height > 0.12 * slide_area:
                    name = f"S{i}_{len(rec['figures']) + 1}.{s.image.ext}"
                    (Path(out_dir) / "figures").mkdir(parents=True, exist_ok=True)
                    (Path(out_dir) / "figures" / name).write_bytes(s.image.blob)
                    rec["figures"].append(name)
            if getattr(s, "has_table", False) and s.has_table:
                rows = [[c.text.strip() for c in row.cells] for row in s.table.rows]
                rec["tables"].append(rows)
            elif getattr(s, "has_text_frame", False) and s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    pr = paragraph_record(p, plain)
                    if pr:
                        rec["items"].append(pr)
        if not rec["title"]:
            big = [it for it in rec["items"] if it["heading"] and it["size"] >= 20]
            if big:
                rec["title"] = big[0]["text"]
        if slide.has_notes_slide:
            rec["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(rec)
    return slides


def to_markdown(slides, name):
    out = [f"# Extracted slides: {name}\n",
           "Bold = emphasized on the slide (bold/underline/highlight/colored text).\n"]
    for s in slides:
        out.append(f"\n## Slide {s['slide']}: {s['title'] or '(no title)'}")
        for it in s["items"]:
            out.append("  " * it["level"] + ("- [heading] " if it["heading"] else "- ") + it["md"])
        for t in s["tables"]:
            out.append("")
            for r, row in enumerate(t):
                out.append("| " + " | ".join(c.replace("\n", " ") for c in row) + " |")
                if r == 0:
                    out.append("|" + "---|" * len(row))
        if s["figures"]:
            out.append(f"\n_[figure(s) saved: {', '.join(s['figures'])}]_")
        elif s["images"]:
            out.append(f"\n_[{s['images']} small image(s)/icons]_")
        if s["notes"]:
            out.append(f"\n> Speaker notes: {s['notes']}")
    return "\n".join(out) + "\n"


def main(pptx_path, out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    slides = extract(pptx_path, out_dir)
    (out_dir / "slides.json").write_text(json.dumps(slides, indent=1, ensure_ascii=False))
    (out_dir / "slides.md").write_text(to_markdown(slides, Path(pptx_path).name))
    print(f"  extracted {len(slides)} slides from {Path(pptx_path).name}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
